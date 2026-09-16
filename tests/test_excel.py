"""Excel 输入层的测试。

## 取材

结构照抄真实材料（某小企业的小企业报表，1993 年行业会计制度格式），
但**数字全部是编的** —— 真实标的的数据不进仓库。

## 两条必须锁住的

**① 左右两栏要切开。** 小企业的资产负债表是 T 型布局，
一行里同时装着资产和负债。不切开的话，标签取自左栏、行次取自右栏，
**混起来而且不报错**。

**② 主数值列必须由表头决定，不能假设顺序。**
资产负债表的列是「年初数 / 期末数」—— **年初在前**。
照搬「取第一个数值列」会把去年的数当成今年的。
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from ingest import excel                                # noqa: E402


def _balance_rows() -> list[list[object]]:
    """左右两栏式资产负债表（结构照真实材料，数字虚构）。"""
    return [
        [None, None, "资产负债表", None, None, None, None, "会工01表"],
        [None, None, None, None, None, None, None, None],
        ["编制单位:某某环境监测有限公司", None, None, "2024 年12 月", "31 日", None, None, "单位：元"],
        ["资    产", "行次", "年初数", "期末数",
         "负债及所有者权益", "行次", "年初数", "期末数"],
        ["流动资产：", None, None, None, "流动负债：", None, None, None],
        ["  货币资金", "1", 21_470_366.19, 24_781_633.06, "  短期借款", "51", None, None],
        ["  短期投资", "2", None, None, "  应付票据", "52", None, None],
        ["  应收票据", "3", 2_395_000.0, 998_812.37, "  应付账款", "53", 8_498_381.02, 6_643_651.85],
        ["  应收账款", "4", 18_878_764.08, 16_247_901.49, "  预收账款", "54", None, None],
        ["    减：坏账准备", "5", None, None, "  其他应付款", "55", 2_722_070.0, 45_720.0],
        ["  应收账款净额", "6", 18_878_764.08, 16_247_901.49, "  应付工资", "56", 4_800_000.0, 5_397_768.60],
        ["  存货", "10", 659_861.49, 514_722.82, "  未交税金", "58", 1_301_261.27, 1_437_373.28],
        ["流动资产合计", "20", 46_501_768.76, 45_286_293.64, "流动负债合计", "70", 20_277_321.29, 25_503_733.68],
        ["固定资产原价", "24", 34_939_808.84, 35_825_021.02, "  实收资本", "91", 11_000_000.0, 11_000_000.0],
        ["    减：累计折旧", "25", 17_732_925.90, 20_949_244.96, "  未分配利润", "95", 29_631_330.41, 20_858_336.02],
        ["  固定资产净值", "26", 17_206_882.94, 14_875_776.06, "所有者权益合计", "96", 43_431_330.41, 34_658_336.02],
        ["资产总计", "45", 63_708_651.70, 60_162_069.70,
         "负债及所有者权益总计", "100", 63_708_651.70, 60_162_069.70],
    ]


def _income_rows() -> list[list[object]]:
    """损益表（旧制度格式）。"""
    return [
        [None, None, None, "损    益    表"],
        [None, None, None, "会工02表"],
        ["编制单位:某某环境监测有限公司", "2024 年", "12 月", "单位：元"],
        [None, None, None, None],
        ["项              目", "行    次", "本  月  数", "本 年 累 计"],
        ["       一、产品销售收入", "1", 9_694_226.70, 79_270_383.88],
        ["           减：产品销售成本", "2", 7_468_869.33, 50_543_383.95],
        ["               产品销售税金及附加", "4", 33_871.60, 393_069.81],
        ["       二、产品销售利润", "5", 2_191_485.77, 28_333_930.12],
        ["           减：管理费用", "7", 2_169_770.21, 18_190_465.39],
        ["               财务费用", "8", -22_421.82, -150_077.24],
        ["       三、营业利润", "9", 44_137.38, 10_293_541.97],
        ["       四、利润总额", "15", 40_737.38, 10_286_764.86],
        ["           减：所得税", "16", 1_184_759.25, 1_184_759.25],
        ["       五、净利润", "17", -1_144_021.87, 9_102_005.61],
    ]


class TestReadWorkbook(unittest.TestCase):
    def test_reads_xlsx(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "t.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "表1"
            for r, row in enumerate(_income_rows(), 1):
                for c, v in enumerate(row, 1):
                    ws.cell(r, c).value = v
            wb.save(p)

            got = excel.read_workbook(p)
            self.assertEqual([s.name for s in got.sheets], ["表1"])
            self.assertGreater(got.sheets[0].n_rows, 10)

    def test_missing_file_says_so(self):
        with self.assertRaises(FileNotFoundError):
            excel.read_workbook("/nonexistent/nope.xls")

    def test_unknown_suffix(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "t.csv"
            p.write_text("a,b\n1,2")
            with self.assertRaises(ValueError) as ctx:
                excel.read_workbook(p)
            self.assertIn(".xls", str(ctx.exception))


class TestLabelColumns(unittest.TestCase):
    def test_two_label_columns(self):
        self.assertEqual(excel.label_columns(_balance_rows()), [0, 4])

    def test_one_label_column(self):
        self.assertEqual(excel.label_columns(_income_rows()), [0])


class TestTwoSidedSplit(unittest.TestCase):
    """**这条是最关键的 —— 切不开就会把两栏混起来。**"""

    def test_detects_and_splits(self):
        rows = _balance_rows()
        self.assertEqual(excel.two_sided_split(rows), 4)
        blocks = excel.split_blocks(rows)
        self.assertEqual(len(blocks), 2)

    def test_single_sided_income_is_untouched(self):
        rows = _income_rows()
        self.assertIsNone(excel.two_sided_split(rows))
        self.assertEqual(len(excel.split_blocks(rows)), 1)

    def test_left_block_keeps_its_own_row_numbers(self):
        """**切不开的症状**：货币资金的行次变成短期借款的 51。"""
        blocks = excel.split_blocks(_balance_rows())
        out, _ = excel.to_rows(blocks[0])
        cash = next(r for r in out if r[0].strip() == "货币资金")
        self.assertEqual(cash[1], "1")          # 不是 51

    def test_right_block_has_liabilities(self):
        blocks = excel.split_blocks(_balance_rows())
        out, _ = excel.to_rows(blocks[1])
        labels = [r[0].strip() for r in out]
        self.assertIn("短期借款", labels)
        self.assertIn("负债及所有者权益总计", labels)
        self.assertNotIn("货币资金", labels)


class TestColumnRoles(unittest.TestCase):
    """**主数值列必须由表头决定。**"""

    def test_balance_primary_is_end_of_year_not_beginning(self):
        """资产负债表的列是「年初数 / 期末数」—— 年初在前。

        照搬「取第一个数值列」会把去年的数当成今年的。
        """
        blocks = excel.split_blocks(_balance_rows())
        out, roles = excel.to_rows(blocks[0])
        self.assertEqual(roles.primary_name.strip(), "期末数")
        self.assertEqual(roles.other_name.strip(), "年初数")

        cash = next(r for r in out if r[0].strip() == "货币资金")
        self.assertEqual(cash[2], "24781633.06")   # 期末
        self.assertEqual(cash[3], "21470366.19")   # 年初

    def test_income_primary_is_year_to_date(self):
        out, roles = excel.to_rows(_income_rows())
        self.assertEqual(roles.primary_name.replace(" ", ""), "本年累计")
        self.assertEqual(roles.other_name.replace(" ", ""), "本月数")

        rev = next(r for r in out if "产品销售收入" in r[0])
        self.assertEqual(rev[2], "79270383.88")    # 本年累计
        self.assertEqual(rev[3], "9694226.7")      # 本月数

    def test_note_column_not_taken_from_other_side(self):
        blocks = excel.split_blocks(_balance_rows())
        _, roles = excel.to_rows(blocks[0])
        self.assertEqual(roles.note, 1)


class TestToRows(unittest.TestCase):
    def test_skips_preamble_rows(self):
        out, _ = excel.to_rows(_income_rows())
        labels = [r[0] for r in out]
        self.assertNotIn("损    益    表", labels)
        self.assertNotIn("会工02表", labels)
        self.assertIn("一、产品销售收入", labels)

    def test_label_only_rows_kept(self):
        """段标题（只有科目名没有数字）要留着，下游会跳过。"""
        out, _ = excel.to_rows(_balance_rows())
        section = [r for r in out if r[0].strip() == "流动资产："]
        self.assertEqual(len(section), 1)
        self.assertEqual(section[0][2], "")


class TestClassify(unittest.TestCase):
    def test_balance(self):
        self.assertEqual(excel.classify([]) if False else
                         _classify(_balance_rows()), "balance")

    def test_income(self):
        self.assertEqual(_classify(_income_rows()), "income")

    def test_unknown(self):
        self.assertEqual(_classify([["随便", "写点什么"]]), "unknown")


def _classify(rows):
    from financials.from_excel import classify
    return classify(rows)


class TestIsLabel(unittest.TestCase):
    def test_numbers_are_not_labels(self):
        for t in ("1", "24,781,633.06", "", "  ", "2024 年12 月", "(100)"):
            self.assertFalse(excel._is_label(t), repr(t))

    def test_account_names_are_labels(self):
        for t in ("货币资金", "资产总计", "  短期借款"):
            self.assertTrue(excel._is_label(t), t)


if __name__ == "__main__":
    unittest.main()
